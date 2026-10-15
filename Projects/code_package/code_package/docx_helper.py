"""docx_helper.py -- tiny paper-layout helper on top of python-docx."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class Paper:
    def __init__(self, title, keywords, running="", refdict=None):
        self.doc = Document()
        self.refdict = refdict or {}
        self.cites = []
        self.nfig = 0
        self.ntab = 0
        self.sec = 0
        self.sub = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"; st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.12
        for lvl, sz in ((1, 13), (2, 11.5), (3, 11)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"; h.font.size = Pt(sz); h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0); h.font.italic = (lvl == 3)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            h.paragraph_format.space_before = Pt(10 if lvl == 1 else 6)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        # title block
        t = self.doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run(title); r.bold = True; r.font.size = Pt(17)
        t.paragraph_format.space_after = Pt(6)
        a = self.doc.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = a.add_run("[Author names]  |  [Department, Institution]  |  [Corresponding e-mail]")
        r.italic = True; r.font.size = Pt(10)
        self.keywords = keywords
        self._footer(running)

    def _footer(self, running):
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run((running + "   |   Page ") if running else "Page ")
        r.font.size = Pt(8.5)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            run = p.add_run(); run.font.size = Pt(8.5)
            if tag:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)

    # ---- citations:  [@key1,key2] -> [n1,n2] in order of first appearance
    def _c(self, text):
        def sub(m):
            nums = []
            for k in m.group(1).split(","):
                k = k.strip()
                if k not in self.cites: self.cites.append(k)
                nums.append(self.cites.index(k) + 1)
            return "[" + ",".join(str(n) for n in sorted(nums)) + "]"
        return re.sub(r"\[@([A-Za-z0-9_,\s]+)\]", sub, text)

    # ---- text
    def abstract(self, text):
        text = self._c(text)
        p = self.doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.right_indent = Cm(0.8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run("Abstract. "); r.bold = True; r.font.size = Pt(10)
        self._runs(p, text)
        for r in p.runs[1:]: r.font.size = Pt(10)
        k = self.doc.add_paragraph(); k.paragraph_format.left_indent = Cm(0.8)
        r = k.add_run("Keywords: "); r.bold = True; r.font.size = Pt(10)
        r = k.add_run(self.keywords); r.italic = True; r.font.size = Pt(10)

    def h1(self, t):
        if not t.startswith(("Appendix", "References")):
            self.sec += 1; self.sub = 0; t = f"{self.sec}. {t}"
        self.doc.add_heading(t, level=1)
    def h2(self, t):
        self.sub += 1; self.doc.add_heading(f"{self.sec}.{self.sub} {t}", level=2)
    def h3(self, t): self.doc.add_heading(t, level=3)

    def p(self, text, italic=False, indent=False):
        text = self._c(text)
        para = self.doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._runs(para, text, italic)
        if indent: para.paragraph_format.left_indent = Cm(0.6)
        return para

    def _runs(self, para, text, italic=False):
        # **bold** inline markup
        parts = text.split("**")
        for i, seg in enumerate(parts):
            if not seg: continue
            r = para.add_run(seg); r.bold = (i % 2 == 1); r.italic = italic

    def bullets(self, items):
        for it in items:
            it = self._c(it)
            para = self.doc.add_paragraph(style="List Bullet"); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, it)

    def numbered(self, items):
        for it in items:
            it = self._c(it)
            para = self.doc.add_paragraph(style="List Number"); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, it)

    def eq(self, text, num=None):
        para = self.doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(text); r.font.name = "Cambria Math"; r.font.size = Pt(11)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
        if num: para.add_run(f"        ({num})")
        para.paragraph_format.space_after = Pt(4)

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            para = self.doc.add_paragraph(); para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.left_indent = Cm(0.7)
            r = para.add_run(line if line else " "); r.font.name = "Courier New"; r.font.size = Pt(8.5)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- figures and tables
    def fig(self, path, caption, width=13.5):
        caption = self._c(caption)
        self.nfig += 1
        para = self.doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True; para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(path, width=Cm(width))
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = c.add_run(f"Figure {self.nfig}. "); r.bold = True; r.font.size = Pt(9.5)
        r = c.add_run(caption); r.font.size = Pt(9.5)
        c.paragraph_format.space_after = Pt(8)
        return self.nfig

    def table(self, rows, caption, widths=None, font=9):
        caption = self._c(caption)
        self.ntab += 1
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.keep_with_next = True; c.paragraph_format.space_after = Pt(2)
        r = c.add_run(f"Table {self.ntab}. "); r.bold = True; r.font.size = Pt(9.5)
        r = c.add_run(caption); r.font.size = Pt(9.5)
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j); cell.text = ""
                para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(0)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(str(val)); run.font.size = Pt(font); run.bold = (i == 0)
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E7E6E6"); tcPr.append(shd)
        if widths:
            for j, w in enumerate(widths):
                for row in t.rows: row.cells[j].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return self.ntab

    def refs(self, extra=()):
        self.h1("References")
        for k in extra:
            if k not in self.cites: self.cites.append(k)
        items = [self.refdict[k] for k in self.cites]
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph(); para.paragraph_format.left_indent = Cm(0.7)
            para.paragraph_format.first_line_indent = Cm(-0.7); para.paragraph_format.space_after = Pt(2)
            r = para.add_run(f"[{i}] {it}"); r.font.size = Pt(9.5)

    def pagebreak(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
